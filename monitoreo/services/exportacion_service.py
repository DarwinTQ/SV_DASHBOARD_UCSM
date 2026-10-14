"""Exportación de resultados a Excel/CSV para el Capítulo V de la tesis."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from monitoreo.repositories.metricas_repository import MetricasRepository
from monitoreo.services.metricas_service import MetricasService

VERDE_UCSM = "0CB967"


class ExportacionService:
    """Genera el libro de Excel con las hojas de resultados de la tesis."""

    def __init__(self, metricas_service=None, metricas_repository=None):
        self._metricas = metricas_service or MetricasService()
        self._repo = metricas_repository or MetricasRepository()

    def generar_excel(self, desde=None, hasta=None, dispositivo_id=None, tipo=None) -> bytes:
        libro = Workbook()
        datos = self._metricas.resultados_completos(desde, hasta, dispositivo_id, tipo)

        self._hoja_resumen(libro.active, datos)
        self._hoja_tiempos(libro.create_sheet("Tiempos"), desde, hasta, dispositivo_id, tipo)
        self._hoja_matriz_confusion(libro.create_sheet("Matriz de confusión"), datos["matriz_confusion"])
        self._hoja_disponibilidad(libro.create_sheet("Disponibilidad"), datos)
        self._hoja_alertas(libro.create_sheet("Alertas"), datos["alertas"])
        self._hoja_incidencias(libro.create_sheet("Incidencias"), desde, hasta, dispositivo_id, tipo)

        buffer = io.BytesIO()
        libro.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ #
    def _titulo(self, hoja, texto: str, fila: int = 1) -> None:
        celda = hoja.cell(row=fila, column=1, value=texto)
        celda.font = Font(bold=True, size=13, color=VERDE_UCSM)

    def _encabezados(self, hoja, fila: int, columnas: list) -> None:
        for indice, nombre in enumerate(columnas, start=1):
            celda = hoja.cell(row=fila, column=indice, value=nombre)
            celda.font = Font(bold=True)

    def _hoja_resumen(self, hoja, datos: dict) -> None:
        hoja.title = "Resumen"
        self._titulo(hoja, "Resumen general — Sistema de Monitoreo Vehicular UCSM")
        fila = 3
        hoja.cell(row=fila, column=1, value="Tiempo de respuesta total (ms)")
        est = datos["tiempos"]["tiempo_respuesta_total_ms"]
        hoja.cell(row=fila, column=2, value=f"n={est['n']}, media={est['media']}, p95={est['p95']}")
        fila += 1
        hoja.cell(row=fila, column=1, value="Disponibilidad (%)")
        hoja.cell(row=fila, column=2, value=datos["disponibilidad"]["disponibilidad_pct"])
        fila += 1
        hoja.cell(row=fila, column=1, value="Alertas enviadas / total")
        hoja.cell(row=fila, column=2, value=f"{datos['alertas']['enviadas']} / {datos['alertas']['total']}")
        fila += 1
        hoja.cell(row=fila, column=1, value="Precisión / Exhaustividad / F1")
        mc = datos["matriz_confusion"]
        hoja.cell(row=fila, column=2, value=f"{mc['precision']} / {mc['exhaustividad']} / {mc['f1']}")

    def _hoja_tiempos(self, hoja, desde, hasta, dispositivo_id, tipo) -> None:
        self._titulo(hoja, "Tabla 25. Tiempo de respuesta entre el evento y la alerta")
        columnas = [
            "ID incidencia", "Fecha/hora", "Tipo", "Latencia red (ms)", "Tiempo captura (ms)",
            "Tiempo inferencia (ms)", "Tiempo BD (ms)", "Latencia Telegram (ms)", "Tiempo total (ms)", "FPS",
        ]
        self._encabezados(hoja, 3, columnas)
        fila = 4
        for inc in self._repo.incidencias_en_rango(desde, hasta, dispositivo_id, tipo).select_related("tiempos", "tipo"):
            t = getattr(inc, "tiempos", None)
            hoja.append([
                inc.pk, inc.fecha_hora.strftime("%Y-%m-%d %H:%M:%S"), inc.tipo.nombre,
                t.latencia_red_ms if t else None, t.tiempo_captura_ms if t else None,
                t.tiempo_inferencia_ms if t else None, t.tiempo_bd_ms if t else None,
                t.latencia_telegram_ms if t else None, t.tiempo_respuesta_total_ms if t else None,
                t.fps if t else None,
            ])

    def _hoja_matriz_confusion(self, hoja, matriz: dict) -> None:
        self._titulo(hoja, "Matriz de confusión")
        hoja.append([])
        hoja.append(["", "Predicho positivo", "Predicho negativo"])
        hoja.append(["Real positivo", matriz["conteo"]["VP"], matriz["conteo"]["FN"]])
        hoja.append(["Real negativo", matriz["conteo"]["FP"], matriz["conteo"]["VN"]])
        hoja.append([])
        hoja.append(["Precisión", matriz["precision"]])
        hoja.append(["Exhaustividad (Recall)", matriz["exhaustividad"]])
        hoja.append(["F1", matriz["f1"]])
        hoja.append(["Exactitud", matriz["exactitud"]])
        hoja.append(["Incidencias sin validar", matriz["sin_validar"]])

    def _hoja_disponibilidad(self, hoja, datos: dict) -> None:
        self._titulo(hoja, "Tabla 26. Disponibilidad del sistema y alertas generadas")
        disp = datos["disponibilidad"]
        hoja.append([])
        hoja.append(["Total de muestras", disp["total_muestras"]])
        hoja.append(["Muestras en línea", disp["muestras_en_linea"]])
        hoja.append(["Disponibilidad (%)", disp["disponibilidad_pct"]])

    def _hoja_alertas(self, hoja, alertas: dict) -> None:
        self._titulo(hoja, "Resumen de alertas Telegram")
        hoja.append([])
        hoja.append(["Total", alertas["total"]])
        hoja.append(["Enviadas", alertas["enviadas"]])
        hoja.append(["Fallidas", alertas["fallidas"]])
        hoja.append(["Tasa de éxito (%)", alertas["tasa_exito"]])
        hoja.append(["Intentos promedio", alertas["intentos_promedio"]])

    def _hoja_incidencias(self, hoja, desde, hasta, dispositivo_id, tipo) -> None:
        self._titulo(hoja, "Incidencias")
        columnas = ["ID", "Fecha/hora", "Tipo", "Riesgo", "Dispositivo", "Placa", "Estado", "Validación"]
        self._encabezados(hoja, 3, columnas)
        for inc in self._repo.incidencias_en_rango(desde, hasta, dispositivo_id, tipo).select_related(
            "tipo", "dispositivo__vehiculo",
        ):
            hoja.append([
                inc.pk, inc.fecha_hora.strftime("%Y-%m-%d %H:%M:%S"), inc.tipo.nombre, inc.tipo.nivel_riesgo,
                inc.dispositivo.codigo, inc.dispositivo.vehiculo.placa, inc.estado, inc.validacion or "",
            ])

    # ------------------------------------------------------------------ #
    def generar_csv_incidencias(self, desde=None, hasta=None, dispositivo_id=None, tipo=None) -> str:
        buffer = io.StringIO()
        escritor = csv.writer(buffer)
        escritor.writerow(["ID", "Fecha/hora", "Tipo", "Riesgo", "Dispositivo", "Placa", "Estado", "Validación"])
        for inc in self._repo.incidencias_en_rango(desde, hasta, dispositivo_id, tipo).select_related(
            "tipo", "dispositivo__vehiculo",
        ):
            escritor.writerow([
                inc.pk, inc.fecha_hora, inc.tipo.nombre, inc.tipo.nivel_riesgo,
                inc.dispositivo.codigo, inc.dispositivo.vehiculo.placa, inc.estado, inc.validacion or "",
            ])
        return buffer.getvalue()
